import csv, hashlib, json, re
from pathlib import Path

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'automated_benchmark_results'
BLIND=OUT/'automated_blinded_adjudication.tsv'
MAN=OUT/'adjudication_freeze_manifest.json'

def test_frozen_blind_units_and_no_leakage():
    rows=list(csv.DictReader(BLIND.open(),delimiter='\t'))
    assert len(rows)==689
    assert {r['blinded_unit_id'] for r in rows}=={f'FB{i:04d}' for i in range(1,690)}
    text=BLIND.read_text().lower()
    assert 'phagemine' not in text and 'pharokka' not in text and 'tool_blinding_key' not in text

def test_manifest_checksums_and_method():
    m=json.loads(MAN.read_text())
    assert m['method']=='AUTOMATED_EVIDENCE_ADJUDICATION'
    for fn,h in m['outputs'].items():
        assert hashlib.sha256((OUT/fn).read_bytes()).hexdigest()==h

def test_categories_and_confidence_are_controlled():
    rows=list(csv.DictReader(BLIND.open(),delimiter='\t'))
    allowed={'CORRECT','PARTIALLY_CORRECT','TOO_GENERAL','UNSUPPORTED/INCORRECT','UNRESOLVABLE','NOT_EVALUABLE'}
    assert {r['adjudication_class'] for r in rows} <= allowed
    assert {r['automated_confidence'] for r in rows} <= {'HIGH','MODERATE','LOW'}

def test_workbook_has_no_hidden_sheets_or_formulas():
    from openpyxl import load_workbook
    wb=load_workbook(OUT/'automated_blinded_adjudication.xlsx',data_only=False)
    assert all(s.sheet_state=='visible' for s in wb.worksheets)
    assert all(cell.data_type!='f' for ws in wb.worksheets for row in ws.iter_rows() for cell in row)
